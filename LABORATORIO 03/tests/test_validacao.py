from openpyxl import load_workbook
import pandas as pd
import pytest

from validacao import consenso, kappa, sorteio
from validacao.rotulos import carregar_avaliadores, tabela_da_dimensao


# ---------------------------------------------------------------- utilitários

def escrever_rotulos(pasta, avaliador, repositorios, releases):
    """Grava rotulos_<avaliador>.xlsx com as duas abas, como o avaliador entregaria."""
    with pd.ExcelWriter(pasta / f"rotulos_{avaliador}.xlsx") as arquivo:
        pd.DataFrame(repositorios, columns=["full_name", "tipo_projeto", "entregas_reais"]
                     ).to_excel(arquivo, sheet_name="repositorios", index=False)
        pd.DataFrame(releases, columns=["full_name", "tag_name", "corretiva"]
                     ).to_excel(arquivo, sheet_name="releases", index=False)


def tabela(*linhas):
    """Itens nas linhas, avaliadores A, B, C nas colunas."""
    return pd.DataFrame(list(linhas), columns=["A", "B", "C"]).astype(object)


# ---------------------------------------------------------------- kappa

def test_kappa_concordancia_total_vale_1():
    resultado = kappa.calcular_kappa(tabela(["sim"] * 3, ["nao"] * 3, ["sim"] * 3), ["sim", "nao"])
    assert resultado["kappa"] == pytest.approx(1)
    assert resultado["unanimidade"] == 1


def test_kappa_calculado_a_mao():
    # P_i: unânime = 1; 2-1 = (2² + 1² - 3) / (3·2) = 1/3  ->  P̄ = (1 + 1 + 1/3 + 1/3) / 4 = 2/3
    # p_sim = p_nao = 6/12 = 0,5  ->  P_e = 0,5  ->  κ = (2/3 - 0,5) / (1 - 0,5) = 1/3
    t = tabela(["sim"] * 3, ["nao"] * 3, ["sim", "sim", "nao"], ["nao", "nao", "sim"])
    resultado = kappa.calcular_kappa(t, ["sim", "nao"])
    assert resultado["kappa"] == pytest.approx(1 / 3)
    assert resultado["unanimidade"] == 0.5


def test_kappa_discordancia_sistematica_e_negativo():
    t = tabela(["sim", "nao", "nao"], ["nao", "sim", "sim"], ["sim", "nao", "nao"],
               ["nao", "sim", "sim"])
    assert kappa.calcular_kappa(t, ["sim", "nao"])["kappa"] < 0


def test_kappa_indefinido_com_uma_unica_categoria():
    resultado = kappa.calcular_kappa(tabela(["sim"] * 3, ["sim"] * 3), ["sim", "nao"])
    assert resultado["kappa"] is None
    assert kappa.interpretar(resultado["kappa"]) == "indefinido"


def test_kappa_ignora_itens_incompletos():
    t = tabela(["sim"] * 3, ["nao"] * 3, ["sim", None, "nao"])
    resultado = kappa.calcular_kappa(t, ["sim", "nao"])
    assert resultado["itens"] == 2 and resultado["itens_incompletos"] == 1
    assert resultado["kappa"] == pytest.approx(1)


def test_kappa_sem_itens_completos():
    resultado = kappa.calcular_kappa(tabela(["sim", None, None]), ["sim", "nao"])
    assert resultado["itens"] == 0 and resultado["kappa"] is None


@pytest.mark.parametrize("valor, esperado", [
    (-0.1, "pobre"), (0, "leve"), (0.20, "leve"), (0.21, "razoável"), (0.40, "razoável"),
    (0.41, "moderada"), (0.60, "moderada"), (0.61, "substancial"), (0.80, "substancial"),
    (0.81, "quase perfeita"), (1, "quase perfeita"),
])
def test_interpretacao_landis_koch(valor, esperado):
    assert kappa.interpretar(valor) == esperado


# ---------------------------------------------------------------- leitura das planilhas

@pytest.fixture
def pasta_rotulos(tmp_path):
    repos = {"A": [["o/a", "ferramenta CLI", "sim"], ["o/b", "outro", "nao"]],
             "B": [["o/a", "ferramenta cli ", "sim"], ["o/b", "outro", "incerto"]],
             "C": [["o/a", "biblioteca/framework", "sim"], ["o/b", "outro", None]]}
    rels = {"A": [["o/a", "v1", "sim"], ["o/a", "v2", "nao"]],
            "B": [["o/a", "v1", "sim"], ["o/a", "v2", "sim"]],
            "C": [["o/a", "v1", "nao"], ["o/a", "v2", "nao"]]}
    for avaliador in "ABC":
        escrever_rotulos(tmp_path, avaliador, repos[avaliador], rels[avaliador])
    (tmp_path / "~$rotulos_A.xlsx").write_bytes(b"")  # bloqueio do Excel: deve ser ignorado
    return tmp_path


def test_leitura_normaliza_caixa_e_espacos(pasta_rotulos):
    avaliadores = carregar_avaliadores(pasta_rotulos)
    assert list(avaliadores) == ["A", "B", "C"]
    t = tabela_da_dimensao(avaliadores, "tipo_projeto")
    assert t.loc["o/a"].tolist() == ["ferramenta CLI", "ferramenta CLI", "biblioteca/framework"]
    assert tabela_da_dimensao(avaliadores, "entregas_reais").loc["o/b", "C"] is None


def test_leitura_rejeita_rotulo_invalido(tmp_path):
    escrever_rotulos(tmp_path, "A", [["o/a", "app", "sim"]], [["o/a", "v1", "sim"]])
    with pytest.raises(ValueError, match="app"):
        tabela_da_dimensao(carregar_avaliadores(tmp_path), "tipo_projeto")


def test_kappa_main_grava_resultado(pasta_rotulos, capsys):
    kappa.main(["--pasta", str(pasta_rotulos)])
    resultado = pd.read_csv(pasta_rotulos / "resultado_kappa.csv")
    assert resultado["dimensao"].tolist() == ["tipo_projeto", "entregas_reais", "corretiva"]
    assert "A, B, C" in capsys.readouterr().out


def test_main_exige_dois_avaliadores(tmp_path):
    with pytest.raises(SystemExit):
        kappa.main(["--pasta", str(tmp_path)])
    with pytest.raises(SystemExit):
        consenso.main(["--pasta", str(tmp_path)])


# ---------------------------------------------------------------- consenso

@pytest.mark.parametrize("votos, esperado", [
    (["sim", "sim", "nao"], ("sim", "maioria")),
    (["sim", "sim", "sim"], ("sim", "maioria")),
    (["sim", "nao", None], (None, "pendente")),               # 1-1 com um rótulo faltando
    (["sim", "sim", None], ("sim", "maioria")),
    (["outro", "ferramenta CLI", "biblioteca/framework"], (None, "pendente")),  # 1-1-1
    ([None, None, None], (None, "pendente")),
])
def test_votar(votos, esperado):
    assert consenso.votar(votos) == esperado


def test_consenso_lista_empates_e_aplica_decisoes(pasta_rotulos, capsys):
    consenso.main(["--pasta", str(pasta_rotulos)])
    repos = pd.read_csv(pasta_rotulos / "consenso_repositorios.csv")
    assert repos.set_index("full_name").loc["o/a", "tipo_projeto"] == "ferramenta CLI"
    empates = pd.read_csv(pasta_rotulos / "desempate.csv", keep_default_na=False)
    # o/b em entregas_reais: nao, incerto, vazio -> 1-1
    assert empates[["dimensao", "full_name"]].values.tolist() == [["entregas_reais", "o/b"]]
    assert "1 empate(s) sem decisão" in capsys.readouterr().out

    empates.loc[0, "decisao"] = "nao"
    empates.to_csv(pasta_rotulos / "desempate.csv", index=False)
    consenso.main(["--pasta", str(pasta_rotulos)])
    repos = pd.read_csv(pasta_rotulos / "consenso_repositorios.csv").set_index("full_name")
    assert repos.loc["o/b", "entregas_reais"] == "nao"
    assert repos.loc["o/b", "metodo_entregas_reais"] == "reuniao"
    # a decisão é preservada quando a lista de empates é regenerada
    assert pd.read_csv(pasta_rotulos / "desempate.csv")["decisao"].tolist() == ["nao"]
    assert "Todos os itens têm rótulo de consenso" in capsys.readouterr().out

    releases = pd.read_csv(pasta_rotulos / "consenso_releases.csv").set_index("tag_name")
    assert releases["corretiva"].to_dict() == {"v1": "sim", "v2": "nao"}


def test_decisao_invalida_e_rejeitada(pasta_rotulos):
    pd.DataFrame([{"dimensao": "entregas_reais", "full_name": "o/b", "tag_name": "",
                   "votos": "", "decisao": "talvez"}]).to_csv(pasta_rotulos / "desempate.csv",
                                                            index=False)
    with pytest.raises(ValueError, match="talvez"):
        consenso.main(["--pasta", str(pasta_rotulos)])


# ---------------------------------------------------------------- sorteio

def repositorios_ficticios(n):
    return pd.DataFrame({"full_name": [f"o/r{i:03d}" for i in range(n)],
                         "url": [f"https://github.com/o/r{i:03d}" for i in range(n)]})


def test_sorteio_de_repositorios_e_reproduzivel_e_usa_os_100_primeiros():
    base = repositorios_ficticios(100)
    crescida = pd.concat([base, repositorios_ficticios(400).iloc[100:]])  # amostra com 400
    sorteados = sorteio.sortear_repositorios(base)
    assert len(sorteados) == 60 and sorteados["full_name"].is_unique
    assert sorteados.equals(sorteio.sortear_repositorios(crescida))


def test_sorteio_exige_repositorios_suficientes():
    with pytest.raises(ValueError):
        sorteio.sortear_repositorios(repositorios_ficticios(10))


def release(nome, tag, dia, na_janela="True", draft="False", prerelease="False", anterior=""):
    return {"full_name": nome, "tag_name": tag, "published_at": f"2026-01-{dia:02d}T00:00:00Z",
            "url": f"https://github.com/{nome}/releases/tag/{tag}", "na_janela": na_janela,
            "draft": draft, "prerelease": prerelease, "tag_anterior": anterior}


def test_sorteio_de_releases_so_usa_estaveis_da_janela():
    releases = pd.DataFrame(
        [release("o/a", f"v{i}", i + 1) for i in range(8)]
        + [release("o/a", "rc", 20, prerelease="True"), release("o/a", "rascunho", 21, draft="True"),
           release("o/a", "velha", 22, na_janela="False"),
           release("o/b", "v1", 1), release("o/b", "v2", 2), release("o/c", "x", 1)])
    sorteadas = sorteio.sortear_releases(releases, ["o/b", "o/a"])
    por_repo = sorteadas.groupby("full_name")["tag_name"].apply(set)
    assert len(por_repo["o/a"]) == 5 and por_repo["o/a"] <= {f"v{i}" for i in range(8)}
    assert por_repo["o/b"] == {"v1", "v2"}            # menos de 5 elegíveis: entram todas
    assert "o/c" not in por_repo                        # não sorteado
    assert sorteadas.equals(sorteio.sortear_releases(releases, ["o/a", "o/b"]))


@pytest.mark.parametrize("anterior, esperado", [
    ("v1.0", "https://github.com/o/a/compare/v1.0...v1.1"),
    ("n8n@1.0", "https://github.com/o/a/compare/n8n%401.0...v1.1"),
    ("", ""), (float("nan"), ""),
])
def test_link_compare(anterior, esperado):
    assert sorteio.link_compare("o/a", anterior, "v1.1") == esperado


def test_main_do_sorteio_gera_planilha_com_listas_e_links(tmp_path, capsys):
    dados = tmp_path / "dados"
    dados.mkdir()
    repositorios_ficticios(100).to_csv(dados / "repositorios.csv", index=False)
    pd.DataFrame([release(f"o/r{i:03d}", f"v{j}", j + 1, anterior=f"v{j - 1}" if j else "")
                  for i in range(100) for j in range(6)]).to_csv(dados / "releases.csv", index=False)

    sorteio.main(["--dados", str(dados), "--saida", str(tmp_path)])

    assert len(pd.read_csv(tmp_path / "amostra_ouro_repositorios.csv")) == 60
    assert len(pd.read_csv(tmp_path / "amostra_ouro_releases.csv")) == 300
    planilha = load_workbook(tmp_path / "planilha_rotulagem.xlsx")
    assert planilha.sheetnames == ["instrucoes", "repositorios", "releases"]
    repos, rels = planilha["repositorios"], planilha["releases"]
    assert [c.value for c in repos[1]] == ["full_name", "repositorio", "lista_de_releases",
                                           "tipo_projeto", "entregas_reais", "observacoes"]
    assert repos["B2"].hyperlink.target.startswith("https://github.com/o/r")
    listas = {dv.formula1 for dv in repos.data_validations.dataValidation}
    assert '"biblioteca/framework,aplicacao/servico,ferramenta CLI,outro"' in listas
    assert '"sim,nao,incerto"' in listas
    assert {dv.formula1 for dv in rels.data_validations.dataValidation} == {'"sim,nao"'}
    assert rels.max_row == 301
    # A planilha em branco é lida sem erro: nenhum rótulo preenchido ainda.
    copia = tmp_path / "copia"
    copia.mkdir()
    (copia / "rotulos_A.xlsx").write_bytes((tmp_path / "planilha_rotulagem.xlsx").read_bytes())
    vazia = tabela_da_dimensao(carregar_avaliadores(copia), "corretiva")
    assert len(vazia) == 300 and vazia["A"].isna().all()
    assert "60 repositórios e 300 releases" in capsys.readouterr().out
