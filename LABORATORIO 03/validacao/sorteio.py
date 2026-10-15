"""Sorteio da amostra-ouro e geração da planilha de rotulagem (seção 6 do enunciado).

Uso: ``python -m validacao.sorteio``

Sorteia 60 repositórios e, de cada um, 5 releases estáveis publicadas na janela.
Gera ``planilha_rotulagem.xlsx`` com links para o repositório, a release e a
lista de commits; cada avaliador salva uma cópia como ``rotulos_<letra>.xlsx``.

Base do sorteio: os 100 primeiros repositórios de ``dados/repositorios.csv``.
O pipeline avalia os candidatos numa ordem embaralhada com semente fixa e grava
os aprovados nessa ordem, então os 100 primeiros são os mesmos quando a amostra
cresce para 300 ou mais. Assim o sorteio feito no início da Sprint 2 continua
reproduzível depois da coleta completa, e os 60 sorteados são uma amostra
aleatória da amostra final.
"""

import argparse
from pathlib import Path
from urllib.parse import quote

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
import pandas as pd

from validacao.rotulos import ABA_RELEASES, ABA_REPOSITORIOS, DIMENSOES, PASTA, RAIZ


SEMENTE = 42
N_REPOSITORIOS = 60
N_RELEASES = 5
BASE = 100

INSTRUCOES = [
    "Protocolo de rotulagem da amostra-ouro (seção 6 do enunciado do Lab03)",
    "",
    "1. Salve uma cópia desta planilha como rotulos_<sua letra>.xlsx (ex.: rotulos_A.xlsx) "
    "na pasta LABORATORIO 03/validacao.",
    "2. Rotule SOZINHO: não consulte os colegas nem a saída da heurística automática.",
    "3. Aba 'repositorios', para cada repositório:",
    "   - tipo_projeto: biblioteca/framework, aplicacao/servico, ferramenta CLI ou outro;",
    "   - entregas_reais: as releases representam entregas reais ao usuário? sim, nao ou incerto.",
    "4. Aba 'releases', para cada release: ela é CORRETIVA, ou seja, seu objetivo principal é "
    "corrigir um defeito recente? sim ou nao. Leia as release notes (coluna release) e a lista "
    "de commits (coluna commits).",
    "5. Use os valores da lista suspensa de cada célula. Observações são opcionais.",
    "6. Commite apenas o seu arquivo, citando a sua Issue.",
]


def sortear_repositorios(repositorios, n=N_REPOSITORIOS, base=BASE, semente=SEMENTE):
    universo = repositorios.head(base)
    if len(universo) < n:
        raise ValueError(f"Só há {len(universo)} repositórios para sortear {n}")
    return universo.sample(n=n, random_state=semente).sort_values("full_name")


def _verdadeiro(coluna):
    return coluna.astype(str).str.strip().str.lower() == "true"


def releases_elegiveis(releases):
    """Releases estáveis (nem draft, nem pré-release) publicadas na janela."""
    return releases[_verdadeiro(releases["na_janela"]) & ~_verdadeiro(releases["draft"])
                    & ~_verdadeiro(releases["prerelease"])]


def sortear_releases(releases, nomes, n=N_RELEASES, semente=SEMENTE):
    elegiveis = releases_elegiveis(releases)
    sorteadas = []
    for nome in sorted(nomes):
        do_repositorio = elegiveis[elegiveis["full_name"] == nome]
        sorteadas.append(do_repositorio.sample(n=min(n, len(do_repositorio)),
                                               random_state=semente))
    resultado = pd.concat(sorteadas) if sorteadas else elegiveis.head(0)
    return resultado.sort_values(["full_name", "published_at"])


def link_compare(full_name, anterior, tag):
    if not isinstance(anterior, str) or not anterior.strip():
        return ""  # primeira release da história: não há com o que comparar
    return (f"https://github.com/{full_name}/compare/"
            f"{quote(anterior, safe='')}...{quote(tag, safe='')}")


def _escrever_aba(planilha, nome, linhas, colunas_link, colunas_rotulo):
    aba = planilha.create_sheet(nome)
    cabecalho = list(linhas[0]) if linhas else []
    aba.append(cabecalho)
    for linha in linhas:
        aba.append([linha[c] for c in cabecalho])
    for celula in aba[1]:
        celula.font = Font(bold=True)
        celula.fill = PatternFill("solid", fgColor="DDEBF7")
    for indice, coluna in enumerate(cabecalho, 1):
        letra = aba.cell(row=1, column=indice).column_letter
        aba.column_dimensions[letra].width = 22 if coluna in colunas_rotulo else 45
        if coluna in colunas_link:
            for linha in range(2, len(linhas) + 2):
                celula = aba.cell(row=linha, column=indice)
                if celula.value:
                    celula.hyperlink = celula.value
                    celula.style = "Hyperlink"
        if coluna in colunas_rotulo:
            _, categorias = DIMENSOES[coluna]
            lista = DataValidation(type="list", formula1=f'"{",".join(categorias)}"',
                                   allow_blank=True, showErrorMessage=True,
                                   errorTitle="Valor inválido",
                                   error=f"Use: {', '.join(categorias)}")
            aba.add_data_validation(lista)
            lista.add(f"{letra}2:{letra}{max(len(linhas) + 1, 2)}")
            for linha in range(2, len(linhas) + 2):
                aba.cell(row=linha, column=indice).fill = PatternFill("solid", fgColor="FFF2CC")
    aba.freeze_panes = "B2"


def gerar_planilha(repositorios, releases, caminho):
    planilha = Workbook()
    instrucoes = planilha.active
    instrucoes.title = "instrucoes"
    for texto in INSTRUCOES:
        instrucoes.append([texto])
    instrucoes["A1"].font = Font(bold=True)
    instrucoes.column_dimensions["A"].width = 130
    for linha in instrucoes.iter_rows():
        linha[0].alignment = Alignment(wrap_text=True)

    linhas_repositorios = [{
        "full_name": r.full_name,
        "repositorio": r.url,
        "lista_de_releases": f"{r.url}/releases",
        "tipo_projeto": None, "entregas_reais": None, "observacoes": None,
    } for r in repositorios.itertuples()]
    linhas_releases = [{
        "full_name": r.full_name,
        "tag_name": r.tag_name,
        "publicada_em": str(r.published_at)[:10],
        "release": r.url,
        "commits": link_compare(r.full_name, r.tag_anterior, r.tag_name),
        "corretiva": None, "observacoes": None,
    } for r in releases.itertuples()]
    _escrever_aba(planilha, ABA_REPOSITORIOS, linhas_repositorios,
                  {"repositorio", "lista_de_releases"}, {"tipo_projeto", "entregas_reais"})
    _escrever_aba(planilha, ABA_RELEASES, linhas_releases,
                  {"release", "commits"}, {"corretiva"})
    planilha.save(caminho)


def main(argv=None):
    parser = argparse.ArgumentParser(description="Sorteio da amostra-ouro")
    parser.add_argument("--dados", default=RAIZ / "dados")
    parser.add_argument("--saida", default=PASTA)
    args = parser.parse_args(argv)
    dados, saida = Path(args.dados), Path(args.saida)

    repositorios = sortear_repositorios(pd.read_csv(dados / "repositorios.csv"))
    releases = sortear_releases(pd.read_csv(dados / "releases.csv", dtype=str),
                                repositorios["full_name"])
    repositorios.to_csv(saida / "amostra_ouro_repositorios.csv", index=False)
    releases[["full_name", "tag_name", "published_at", "url", "tag_anterior"]].to_csv(
        saida / "amostra_ouro_releases.csv", index=False)
    gerar_planilha(repositorios, releases, saida / "planilha_rotulagem.xlsx")

    poucas = releases.groupby("full_name").size().loc[lambda s: s < N_RELEASES]
    print(f"{len(repositorios)} repositórios e {len(releases)} releases sorteados "
          f"(semente {SEMENTE}, base: {BASE} primeiros de repositorios.csv).")
    if len(poucas):
        print(f"Repositórios com menos de {N_RELEASES} releases elegíveis: {dict(poucas)}")
    print(f"Planilha: {saida / 'planilha_rotulagem.xlsx'}")


if __name__ == "__main__":
    main()
